import openpyxl
import frappe

PATH = "/Users/bryantantonio/Documents/Hicom/Client/IIB/Import/Master Card.xlsx"
TOL = 0.01


def _s(v):
    return "" if v is None else str(v).strip()


def _f(v):
    try:
        return float(v or 0)
    except (ValueError, TypeError):
        return 0.0


def execute(sample=None):
    """Cross-check imported Master Cards against the source workbook.

    Key invariant: the legacy header's PriceSet / TotalSet equal the
    qty-weighted sum of the base-tier component prices. That single check
    validates the qty mapping, the price mapping and the base-tier
    assignment all at once.
    """
    wb = openpyxl.load_workbook(PATH, read_only=True, data_only=True)
    mc = wb["MastCard"]
    header = next(mc.iter_rows(min_row=1, max_row=1, values_only=True))
    idx = {n: i for i, n in enumerate(header)}

    legacy = {}
    for row in mc.iter_rows(min_row=2, values_only=True):
        legacy[_s(row[idx["MCNo"]])] = {
            "CustCode": _s(row[idx["CustCode"]]),
            "Descr": _s(row[idx["Descr"]]),
            "PriceSet": _f(row[idx["PriceSet"]]),
            "TotalSet": _f(row[idx["TotalSet"]]),
            "Obsolete": _s(row[idx["Obsolete"]]),
            "ManfJoin1": _s(row[idx["ManfJoin1"]]),
            "Currcode": _s(row[idx["Currcode"]]),
        }

    names = frappe.get_all("Master Card", pluck="name")
    if sample:
        names = names[:sample]
    print("Imported Master Cards to verify:", len(names))

    checked = 0
    price_mismatch = []
    total_mismatch = []
    field_mismatch = []
    missing_legacy = []
    orphan_tiers = []
    no_price_rows = []

    for name in names:
        src = legacy.get(name)
        if not src:
            missing_legacy.append(name)
            continue

        doc = frappe.get_doc("Master Card", name)
        checked += 1

        qty_by_comp = {i.component: i.qty for i in doc.items}

        # Every price row must point at a real MOQ child row (idx is 1-based),
        # otherwise it is invisible in the Item Price Set editor.
        moq_idxs = {m.idx for m in doc.moq_items}
        bad = [p.moq_idx for p in doc.price_items if p.moq_idx not in moq_idxs]
        if bad:
            orphan_tiers.append((name, sorted(set(bad)), sorted(moq_idxs)))
        if not doc.price_items:
            no_price_rows.append(name)

        # The lowest tier carries the base price that the legacy header totals
        # were computed from.
        lowest = min(moq_idxs) if moq_idxs else None
        base = [p for p in doc.price_items if p.moq_idx == lowest]

        calc_price = sum(p.material * qty_by_comp.get(p.component, 0) for p in base)
        calc_total = sum(p.total * qty_by_comp.get(p.component, 0) for p in base)

        if abs(calc_price - src["PriceSet"]) > TOL:
            price_mismatch.append((name, round(calc_price, 4), src["PriceSet"]))
        if abs(calc_total - src["TotalSet"]) > TOL:
            total_mismatch.append((name, round(calc_total, 4), src["TotalSet"]))

        expected_customer = f"C{src['CustCode']}" if src["CustCode"] else None
        expected_disabled = 1 if src["Obsolete"] == "-1" else 0
        if doc.customer != expected_customer:
            field_mismatch.append((name, "customer", doc.customer, expected_customer))
        if doc.description != src["Descr"]:
            field_mismatch.append((name, "description", doc.description, src["Descr"]))
        if int(doc.disabled or 0) != expected_disabled:
            field_mismatch.append((name, "disabled", doc.disabled, expected_disabled))
        if _s(doc.custom_manufacturing_joint_1) != src["ManfJoin1"]:
            field_mismatch.append((name, "joint1", doc.custom_manufacturing_joint_1, src["ManfJoin1"]))

    print()
    print(f"Checked: {checked}")
    print(f"PriceSet mismatches (>{TOL}): {len(price_mismatch)}")
    for m in price_mismatch[:15]:
        print("   MC", m[0], "calc:", m[1], "legacy:", m[2])
    print(f"TotalSet mismatches (>{TOL}): {len(total_mismatch)}")
    for m in total_mismatch[:15]:
        print("   MC", m[0], "calc:", m[1], "legacy:", m[2])
    print(f"Header field mismatches: {len(field_mismatch)}")
    for m in field_mismatch[:15]:
        print("   MC", m[0], m[1], "got:", repr(m[2]), "expected:", repr(m[3]))
    print(f"Imported but absent from workbook: {len(missing_legacy)}")
    print(f"Cards with ORPHANED price rows (invisible in UI): {len(orphan_tiers)}")
    for m in orphan_tiers[:15]:
        print("   MC", m[0], "price moq_idx:", m[1], "available MOQ idx:", m[2])
    print(f"Cards with NO price rows at all: {len(no_price_rows)}")
    if no_price_rows[:15]:
        print("   e.g.", no_price_rows[:15])

    return {
        "checked": checked,
        "price_mismatch": len(price_mismatch),
        "total_mismatch": len(total_mismatch),
        "field_mismatch": len(field_mismatch),
        "orphan_tiers": len(orphan_tiers),
        "no_price_rows": len(no_price_rows),
    }
